"""
import_customers_from_docx.py
------------------------------
Reads every Word (.docx) file in a folder, uses the Groq API to pull the
customer's name and email out of each document's text, and saves them into
the existing `customers` table in Supabase (via the app's own database.py /
models.py — see CLAUDE.md's note that this project already talks to
Supabase through SQLAlchemy).

Configuration (read from .env):
    CUSTOMER_DOCS_DIR   Folder containing the .docx files (see .env.example).
    GROQ_API_KEY        API key for the Groq LLM used to extract name/email.
    DATABASE_URL        Already used by database.py.

Run it with:
    python import_customers_from_docx.py

Any missing third-party packages (python-docx, groq) are installed
automatically on first run.
"""

import importlib
import json
import os
import subprocess
import sys

# ---------------------------------------------------------------------------
# 1. Make sure the packages we need are installed before we try to import
#    them. python-docx and groq are not part of the base project, so a fresh
#    checkout won't have them yet.
# ---------------------------------------------------------------------------
_REQUIRED_PACKAGES = {
    # import name -> pip package name
    "docx": "python-docx",
    "groq": "groq",
    "dotenv": "python-dotenv",
}


def _ensure_packages_installed() -> None:
    missing = []
    for import_name, pip_name in _REQUIRED_PACKAGES.items():
        try:
            importlib.import_module(import_name)
        except ImportError:
            missing.append(pip_name)

    if not missing:
        return

    print(f"Installing missing packages: {', '.join(missing)} ...")
    subprocess.run(
        [sys.executable, "-m", "pip", "install", "-q", *missing],
        check=True,
    )
    importlib.invalidate_caches()


_ensure_packages_installed()

import docx  # noqa: E402  (imported after the auto-install step above)
from dotenv import load_dotenv  # noqa: E402
from groq import Groq  # noqa: E402

load_dotenv()

from database import SessionLocal  # noqa: E402
from models import Customer  # noqa: E402
from repository.customers_repository import (  # noqa: E402
    add_customer,
    get_customer_by_email,
    save,
)

GROQ_MODEL = "openai/gpt-oss-120b"


def read_docx_text(path: str) -> str:
    """Flatten a .docx file's paragraphs and tables into plain text."""
    document = docx.Document(path)
    lines = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(": ".join(cells))

    return "\n".join(lines)


def extract_name_and_email(client: Groq, text: str) -> dict:
    """Ask Groq to find the customer's name and email inside `text`."""
    prompt = (
        "The text below comes from a customer intake form. It may describe "
        "an individual or a company with a main contact person.\n\n"
        "Return ONLY a JSON object with exactly two keys: \"name\" and "
        "\"email\". \"name\" is the customer's own name if they are an "
        "individual, or the main contact person's name if it is a business "
        "form. \"email\" is that same person's email address. If either "
        "value cannot be found, use an empty string for it.\n\n"
        f"Text:\n{text}"
    )

    response = client.chat.completions.create(
        model=GROQ_MODEL,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
        temperature=0,
    )
    return json.loads(response.choices[0].message.content)


def save_customer(db, name: str, email: str) -> str:
    """Insert or update the customer row. Returns 'created' or 'updated'."""
    existing = get_customer_by_email(db, email)
    if existing:
        existing.full_name = name
        save(db, existing)
        return "updated"

    add_customer(db, Customer(email=email, full_name=name))
    return "created"


def main() -> None:
    docs_dir = os.getenv("CUSTOMER_DOCS_DIR")
    if not docs_dir:
        sys.exit("CUSTOMER_DOCS_DIR is not set in .env")

    groq_api_key = os.getenv("GROQ_API_KEY")
    if not groq_api_key:
        sys.exit("GROQ_API_KEY is not set in .env")

    if not os.path.isdir(docs_dir):
        sys.exit(f"CUSTOMER_DOCS_DIR '{docs_dir}' is not a folder")

    docx_files = sorted(
        f
        for f in os.listdir(docs_dir)
        if f.lower().endswith(".docx") and not f.startswith("~$")
    )
    if not docx_files:
        sys.exit(f"No .docx files found in '{docs_dir}'")

    client = Groq(api_key=groq_api_key)
    db = SessionLocal()
    results = []

    try:
        for filename in docx_files:
            path = os.path.join(docs_dir, filename)
            try:
                text = read_docx_text(path)
                extracted = extract_name_and_email(client, text)
                name = (extracted.get("name") or "").strip()
                email = (extracted.get("email") or "").strip()

                if not name or not email:
                    results.append((filename, name, email, "skipped (missing name/email)"))
                    continue

                status = save_customer(db, name, email)
                results.append((filename, name, email, status))
            except Exception as exc:  # keep going even if one file fails
                results.append((filename, "", "", f"error: {exc}"))
    finally:
        db.close()

    print("\nResults")
    print("-------")
    for filename, name, email, status in results:
        print(f"{filename}: name={name!r} email={email!r} -> {status}")

    ok = sum(1 for r in results if r[3] in ("created", "updated"))
    print(f"\n{ok}/{len(results)} file(s) saved to the customers table.")


if __name__ == "__main__":
    main()
